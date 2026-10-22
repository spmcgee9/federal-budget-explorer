#!/usr/bin/env python3
"""
Build state-level data (what each state pays in federal taxes and receives in
federal spending) from the New York State Comptroller's annual
"Balance of Payments in the Federal Budget" Excel workbook.

Usage:
    python3 pipeline/build_states.py --ffy 2024 --xlsx raw/states/osc-fedbudget-fy2024.xlsx
    python3 pipeline/build_states.py --ffy 2025 --url <link to the new year's Excel file>

Requires: pip install openpyxl

Writes web/data/states/ffy<YEAR>.json and adds the year to web/data/manifest.json.
The script stops if a state's categories don't add up to its totals, or if its
balance doesn't match the Comptroller's published balance, so a changed layout
fails loudly instead of producing wrong numbers.
"""
import argparse, datetime, json, os, sys, urllib.request
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOL = 0.5  # $0.5 million

STATES = ["Alabama", "Alaska", "Arizona", "Arkansas", "California", "Colorado", "Connecticut",
          "Delaware", "District of Columbia", "Florida", "Georgia", "Hawaii", "Idaho", "Illinois",
          "Indiana", "Iowa", "Kansas", "Kentucky", "Louisiana", "Maine", "Maryland",
          "Massachusetts", "Michigan", "Minnesota", "Mississippi", "Missouri", "Montana",
          "Nebraska", "Nevada", "New Hampshire", "New Jersey", "New Mexico", "New York",
          "North Carolina", "North Dakota", "Ohio", "Oklahoma", "Oregon", "Pennsylvania",
          "Rhode Island", "South Carolina", "South Dakota", "Tennessee", "Texas", "Utah",
          "Vermont", "Virginia", "Washington", "West Virginia", "Wisconsin", "Wyoming"]

# Friendlier names for the Comptroller's column headings
NAMES = {
    "Individual Income Tax Less Refunds": "Individual income tax",
    "Social Insurance": "Payroll taxes (Social Security, Medicare, unemployment)",
    "Corporate Tax": "Corporate income tax",
    "Excise Taxes": "Excise taxes",
    "Estate and Gift Tax": "Estate and gift tax",
    "Total Social Security": "Social Security",
    "Total Medicare": "Medicare",
    "Veterans Benefits": "Veterans' benefits",
    "Federal Employee Benefits": "Federal retiree benefits",
    "Federal Direct Student Loan Costs": "Student loan costs",
    "SNAP": "SNAP (food assistance)",
    "SSI": "Supplemental Security Income",
    "Refundable EITC": "Earned Income Tax Credit (refundable)",
    "Unemployment Compensation": "Unemployment compensation",
    "Federal Pell Grants": "Pell grants",
    "Refundable CTC": "Child Tax Credit (refundable)",
    "Federal Aid for Highways": "Highways",
    "Housing [1]": "Housing assistance",
    "Child Nutrition [2]": "Child nutrition (school meals)",
    "Children's Health Insurance Program": "Children's Health Insurance Program (CHIP)",
    "Elementary and Secondary Education Act (ESEA: Title 1)": "K-12 education (Title I)",
    "Temporary Assistance for Needy Families (TANF)": "TANF (cash assistance)",
    "Special Education": "Special education",
    "Child and Family Services [3]": "Child and family services (incl. Head Start)",
    "Transit Formula Grants": "Transit formula grants",
    "Education Stabilization Fund": "Education Stabilization Fund (pandemic aid)",
    "Transit Infrastructure Grants": "Transit infrastructure grants",
    "Highway Infrastructure Programs": "Highway infrastructure programs",
    "Defense": "Defense contracts",
    "All Other": "Non-defense contracts",
    "Civilian": "Civilian employees",
    "Military": "Military personnel",
    "EIDL Loans": "Small-business disaster loans (EIDL)",
}


def norm(s):
    return " ".join(str(s).split())


def sheet(wb, prefix):
    for ws in wb.worksheets:  # exact title first, then prefix
        if norm(ws.title) == prefix:
            return ws
    for ws in wb.worksheets:
        if norm(ws.title).startswith(prefix):
            return ws
    sys.exit(f"Sheet starting with '{prefix}' not found; the workbook layout has changed")


def rows_by_state(ws):
    out = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        a = norm(a)
        if a in STATES:
            out.setdefault(a, r)
        elif a.startswith("National") or a.startswith("Total: 50"):
            out.setdefault("_national", r)
    missing = [s for s in STATES if s not in out]
    if missing or "_national" not in out:
        sys.exit(f"{ws.title}: missing rows for {missing or 'national total'}")
    return out


POP = {}


def check_pop(ws, col, rows):
    """total / per-capita must equal the row's own population; catches rows whose
    numbers were pasted against the wrong state."""
    if not POP:
        return
    for k, r in rows.items():
        t, pc = ws.cell(r, col).value, ws.cell(r, col + 1).value
        if k == "_national" or not isinstance(t, (int, float)) or not pc:
            continue
        if abs(t * 1e6 / pc / POP[k] - 1) > 0.002:
            sys.exit(f"{ws.title}: row '{k}' holds another state's numbers (total/per-capita "
                     f"implies population {t*1e6/pc:,.0f}, actual {POP[k]:,.0f})")


def single(ws):
    """Sheets with one value per state: the first number to the right of the name."""
    rows = rows_by_state(ws)
    def first_num(r):
        for c in range(2, 6):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                return float(v)
        sys.exit(f"{ws.title}: no number in row {r}")
    if ws.cell(3, 2).value or ws.cell(2, 2).value:
        check_pop(ws, 2, rows)
    return {k: first_num(r) for k, r in rows.items()}


def groups(ws):
    """Sheets with several column groups: name in row 2, 'Total (millions)' below it."""
    rows = rows_by_state(ws)
    cols = [(c.column, norm(c.value)) for c in ws[2] if c.value is not None]
    out = {}
    for col, name in cols:
        if "Total" not in str(ws.cell(3, col).value):
            sys.exit(f"{ws.title}: expected a Total column under '{name}'")
        check_pop(ws, col, rows)
        out[NAMES.get(name, name)] = {k: float(ws.cell(r, col).value or 0) for k, r in rows.items()}
    return out


def node(name, v, children=None):
    n = {"n": name, "v": round(v * 1e6)}
    if children:
        n["c"] = sorted(children, key=lambda c: -c["v"])
    return n


def with_other(name, total, parts, other_name, key):
    """Category with listed programs plus a remainder line."""
    kids = [node(p, vals[key]) for p, vals in parts.items() if vals[key] > 0]
    rest = total[key] - sum(vals[key] for vals in parts.values())
    if rest < -TOL:
        sys.exit(f"{name} programs exceed the category total for {key}")
    if rest > TOL:
        kids.append(node(other_name, rest))
    return node(name, total[key], kids)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ffy", type=int, required=True, help="federal fiscal year the report covers")
    ap.add_argument("--xlsx", help="path to the Comptroller's Excel file")
    ap.add_argument("--url", help="download the Excel file from this address instead")
    ap.add_argument("--report", default=None, help="link to the report PDF (for the page's source links)")
    a = ap.parse_args()

    path = a.xlsx or os.path.join(ROOT, "raw", "states", f"osc-fedbudget-fy{a.ffy}.xlsx")
    if a.url:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        urllib.request.urlretrieve(a.url, path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)

    POP.update({k: v for k, v in single(sheet(wb, "Population")).items()})
    taxes_total = single(sheet(wb, "Taxes Paid"))
    tax_cats = groups(sheet(wb, "Taxes Paid by Category"))
    exp_total = single(sheet(wb, "Expenditures"))
    dp = single(sheet(wb, "Direct Payments"))
    dp_prog = groups(sheet(wb, "Major Direct Payment"))
    gr = single(sheet(wb, "Grants"))
    gr_prog = groups(sheet(wb, "Major Grant"))
    pr = single(sheet(wb, "Procurement (Contracts)"))
    pr_ag = groups(sheet(wb, "Procurement by Major"))
    wg = single(sheet(wb, "Wages and Salaries"))
    oth = single(sheet(wb, "Other Identified Federal"))
    # Not used: "Wages and Salaries by Type" (civilian column fails the population check in
    # the FFY2024 workbook and military exceeds total wages in some states) and
    # "Other Identified Programs" (one program, EIDL, rows shifted; the total sheet is used).
    bop = single(sheet(wb, "Balance of Payments"))
    pop = single(sheet(wb, "Population"))

    keys = STATES + ["_national"]
    out = {}
    for k in keys:
        paid = node("Federal taxes paid", taxes_total[k],
                    [node(n, vals[k]) for n, vals in tax_cats.items()])
        received = node("Federal spending received", exp_total[k], [
            with_other("Direct payments to people", dp, dp_prog, "Other direct payments", k),
            with_other("Grants to state and local governments", gr, gr_prog, "Other grants", k),
            with_other("Federal contracts", pr, pr_ag, "Other contracts", k),
            node("Federal employee wages", wg[k]),
            node("Small-business disaster loans (EIDL)", oth[k]),
        ])
        # checks
        for t, label in ((paid, "taxes"), (received, "spending")):
            s = sum(c["v"] for c in t["c"]) / 1e6
            if abs(s - t["v"] / 1e6) > TOL:
                sys.exit(f"{k}: {label} categories sum to {s:,.1f}M, total is {t['v']/1e6:,.1f}M")
        bal = exp_total[k] - taxes_total[k]
        if abs(bal - bop[k]) > TOL:
            sys.exit(f"{k}: balance {bal:,.1f}M differs from published {bop[k]:,.1f}M")
        out[k] = {"pop": int(pop[k]), "paid": paid, "received": received}

    nat = out.pop("_national")
    meta = {
        "ffy": a.ffy,
        "generated": datetime.date.today().isoformat(),
        "source": f"Office of the New York State Comptroller, New York's Balance of Payments "
                  f"in the Federal Budget, Federal Fiscal Year {a.ffy}",
        "links": {
            "report": a.report or f"https://www.osc.ny.gov/files/reports/pdf/balance-of-payments-ffy-{a.ffy}.pdf",
            "data": a.url or f"https://www.osc.ny.gov/files/reports/excel/fedbudgetfy{a.ffy}data.xlsx",
        },
        "omitted": "The Comptroller's civilian/military wage split is not shown: in the published "
                    "workbook it does not reconcile to the wage totals.",
        "note": "Covers the 50 states and D.C. Net interest on the debt, customs duties, overseas "
                "spending and other amounts that cannot be assigned to a state are excluded.",
    }
    outdir = os.path.join(ROOT, "web", "data", "states")
    os.makedirs(outdir, exist_ok=True)
    json.dump({"meta": meta, "national": nat, "states": out},
              open(os.path.join(outdir, f"ffy{a.ffy}.json"), "w"), separators=(",", ":"))

    mpath = os.path.join(ROOT, "web", "data", "manifest.json")
    manifest = json.load(open(mpath)) if os.path.exists(mpath) else {"years": []}
    years = {y["fy"]: y for y in manifest["years"]}
    years[a.ffy] = {**years.get(a.ffy, {"fy": a.ffy}), "states": True}
    manifest["years"] = sorted(years.values(), key=lambda y: -y["fy"])
    json.dump(manifest, open(mpath, "w"), indent=2)

    print(f"FFY{a.ffy}: taxes ${nat['paid']['v']/1e9:,.1f}B, spending ${nat['received']['v']/1e9:,.1f}B "
          f"across 50 states + DC; all categories and balances check out.")


if __name__ == "__main__":
    main()
